#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
rag_qa/edu_document_loaders/edu_docloader.py

基于 python-docx + RapidOCR 的 Word 文档加载器。
支持纯文本 DOCX 和图片（内嵌图片 OCR）。
"""

import os
import tempfile
from pathlib import Path
from typing import Iterator
from docx import Document as DocxDocument
from langchain_core.documents import Document
from .edu_ocr import get_ocr
from loguru import logger


class OCRDOCLoader:
    """Word 文档加载器：文本段落直接提取，内嵌图片使用 OCR。"""

    def __init__(self, file_path: str):
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"DOCX 文件不存在: {self.file_path}")

    def load(self) -> list[Document]:
        docs = []
        ocr = get_ocr()
        doc = DocxDocument(str(self.file_path))

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                docs.append(Document(
                    page_content=text,
                    metadata={
                        "source": str(self.file_path),
                        "type": "text",
                    },
                ))

        # OCR 内嵌图片
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                try:
                    image_bytes = rel.target_part.blob
                    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                        tmp_path = tmp.name
                        tmp.write(image_bytes)
                    result, _ = ocr(tmp_path)
                    os.unlink(tmp_path)
                    if result:
                        ocr_text = "\n".join([line[0] for line in result])
                        if ocr_text.strip():
                            docs.append(Document(
                                page_content=ocr_text,
                                metadata={
                                    "source": str(self.file_path),
                                    "type": "ocr_image",
                                },
                            ))
                except Exception as e:
                    logger.warning(f"DOCX 图片 OCR 失败: {e}")

        logger.info(f"DOCX 加载完成: {self.file_path}, 共 {len(docs)} 段")
        return docs
